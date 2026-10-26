"""
# // ==============================================================================
# // KCC CLASSSTAMP - WEEK 5 BACKLOG SWEEP TESTS
# // ==============================================================================
# // Covers:
# //   - A-06: 1-click end-of-year batch student promotion
# //   - A-13: Student TOTP QR secret reset endpoint
# //   - S-09: Excuse slip status tracker on the student dashboard
# //   - S-06: Filterable meeting-by-meeting attendance log
# //   - A-10: CSV import 3-step preview (upload -> preview -> confirm)
# //   - T-06: Midterm / Finals split Excel export
# // ==============================================================================
"""

from datetime import date, time, timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.models import (
    AcademicYear,
    AttendanceRecord,
    AuditEvent,
    Class,
    ClassSession,
    ClassSlot,
    Department,
    Enrollment,
    ExcuseSlip,
    OfflineBadge,
    Profile,
    Program,
    Subject,
)
from apps.core.views.helpers.dashboard_helpers import _student_dashboard_stats

User = get_user_model()


class BaseWeek5FixtureMixin:
    def _make_department(self):
        return Department.objects.create(code='CCS', name='College of Computer Studies')

    def _make_program(self, years_total=4):
        return Program.objects.create(
            department=self.department, code='BSIT',
            name='BS Information Technology', years_total=years_total,
        )

    def _make_academic_year(self):
        return AcademicYear.objects.create(
            name='2025-2026', starts_on=date(2025, 6, 1), ends_on=date(2026, 3, 31), is_active=True,
        )

    def _make_admin(self, username='w5_admin'):
        user = User.objects.create_user(username=username, password='Admin!Pass123')
        Profile.objects.create(user=user, role=Profile.ROLE_ADMIN)
        return user

    def _make_teacher(self, username='w5_teacher'):
        user = User.objects.create_user(username=username, password='Teacher!Pass123')
        profile = Profile.objects.create(
            user=user, role=Profile.ROLE_TEACHER, department=self.department,
        )
        return user, profile

    def _make_student(self, username, student_number, year_level=1, approved=True,
                      first_name='Ana', last_name='Cruz'):
        user = User.objects.create_user(
            username=username, password='Student!Pass123',
            first_name=first_name, last_name=last_name,
        )
        profile = Profile.objects.create(
            user=user, role=Profile.ROLE_STUDENT, student_number=student_number,
            program=self.program, year_level=year_level,
            approval_status=(
                Profile.APPROVAL_APPROVED if approved else Profile.APPROVAL_PENDING
            ),
            qr_secret='a' * 64,
        )
        return user, profile

    def _make_class(self, teacher_profile, code='BSIT-1A', subject=None):
        if subject is None:
            subject = Subject.objects.create(
                code=f'SUBJ-{code}', name=f'Subject for {code}',
            )
        return Class.objects.create(
            academic_year=self.academic_year, class_code=code,
            subject=subject, teacher_profile=teacher_profile, grace_minutes=10,
        )

    def _make_slot(self, class_obj, weekday=0, starts_at=None, ends_at=None):
        return ClassSlot.objects.create(
            class_obj=class_obj, weekday=weekday,
            starts_at=starts_at or time(8, 0), ends_at=ends_at or time(9, 0),
            room='Room 1',
        )

    def _make_session(self, slot, session_date, status=ClassSession.STATUS_CLOSED):
        return ClassSession.objects.create(
            slot=slot, session_date=session_date, status=status,
        )

    def _make_record(self, session, student_profile, status):
        return AttendanceRecord.objects.create(
            session=session, student_profile=student_profile,
            status=status, method=AttendanceRecord.METHOD_SYSTEM,
            stamped_at=timezone.now(),
        )


# // ==============================================================================
# // A-06: END-OF-YEAR BATCH STUDENT PROMOTION
# // ==============================================================================
class PromotionTests(TestCase, BaseWeek5FixtureMixin):
    def setUp(self):
        self.department = self._make_department()
        self.program = self._make_program()
        self.academic_year = self._make_academic_year()
        self.admin = self._make_admin()
        self.teacher_user, _ = self._make_teacher()

        _, self.y1a = self._make_student('w5_p1', '2026-9001', year_level=1)
        _, self.y1b = self._make_student('w5_p2', '2026-9002', year_level=1)
        _, self.y4 = self._make_student('w5_p4', '2026-9003', year_level=4)
        _, self.pending = self._make_student('w5_pp', '2026-9004', year_level=1, approved=False)

    def test_promote_all_advances_approved_students(self):
        self.client.force_login(self.admin)
        res = self.client.post(reverse('student_promote_all'))

        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()['status'], 'success')
        self.assertEqual(res.json()['promoted_count'], 2)

        self.y1a.refresh_from_db()
        self.y1b.refresh_from_db()
        self.assertEqual(self.y1a.year_level, 2)
        self.assertEqual(self.y1b.year_level, 2)

        audit = AuditEvent.objects.filter(action='PROMOTE_STUDENTS')
        self.assertTrue(audit.exists())
        self.assertEqual(audit.first().actor, self.admin)

    def test_promotion_capped_at_program_duration(self):
        self.client.force_login(self.admin)
        self.client.post(reverse('student_promote_all'))

        self.y4.refresh_from_db()
        self.assertEqual(self.y4.year_level, 4)

    def test_pending_students_are_not_promoted(self):
        self.client.force_login(self.admin)
        self.client.post(reverse('student_promote_all'))

        self.pending.refresh_from_db()
        self.assertEqual(self.pending.year_level, 1)

    def test_promote_all_requires_admin(self):
        self.client.force_login(self.teacher_user)
        res = self.client.post(reverse('student_promote_all'))
        self.assertEqual(res.status_code, 302)


# // ==============================================================================
# // A-13: STUDENT TOTP QR SECRET RESET
# // ==============================================================================
class QrResetTests(TestCase, BaseWeek5FixtureMixin):
    def setUp(self):
        self.department = self._make_department()
        self.program = self._make_program()
        self.academic_year = self._make_academic_year()
        self.admin = self._make_admin()
        self.teacher_user, _ = self._make_teacher()
        _, self.student = self._make_student('w5_qr', '2026-9101')

    def test_qr_reset_rotates_secret_and_audits(self):
        self.client.force_login(self.admin)
        old_secret = self.student.qr_secret

        res = self.client.post(reverse('student_qr_reset_ajax'), {'student_id': self.student.pk})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()['status'], 'success')

        self.student.refresh_from_db()
        self.assertNotEqual(self.student.qr_secret, old_secret)
        self.assertEqual(len(self.student.qr_secret), 64)
        self.assertTrue(
            AuditEvent.objects.filter(
                action='STUDENT_QR_RESET', object_id=str(self.student.pk),
            ).exists()
        )

    def test_qr_reset_invalidates_unconsumed_offline_badges(self):
        import hashlib
        badge = OfflineBadge.objects.create(
            student_profile=self.student,
            token_hash=hashlib.sha256(b'x' * 32).hexdigest(),
            expires_at=timezone.now() + timedelta(hours=48),
        )

        self.client.force_login(self.admin)
        res = self.client.post(reverse('student_qr_reset_ajax'), {'student_id': self.student.pk})
        self.assertEqual(res.status_code, 200)

        self.assertFalse(OfflineBadge.objects.filter(pk=badge.pk).exists())

    def test_qr_reset_unknown_student_returns_404(self):
        self.client.force_login(self.admin)
        res = self.client.post(reverse('student_qr_reset_ajax'), {'student_id': 999999})
        self.assertEqual(res.status_code, 404)

    def test_qr_reset_requires_admin(self):
        self.client.force_login(self.teacher_user)
        res = self.client.post(reverse('student_qr_reset_ajax'), {'student_id': self.student.pk})
        self.assertEqual(res.status_code, 302)


# // ==============================================================================
# // S-09: EXCUSE SLIP STATUS TRACKER ON STUDENT DASHBOARD
# // ==============================================================================
class ExcuseTrackerTests(TestCase, BaseWeek5FixtureMixin):
    def setUp(self):
        self.department = self._make_department()
        self.program = self._make_program()
        self.academic_year = self._make_academic_year()
        self.teacher_user, self.teacher_profile = self._make_teacher('w5_s9_teacher')
        self.student_user, self.student = self._make_student('w5_s9', '2026-9201')

        subject = Subject.objects.create(code='IT101', name='Introduction to Computing')
        self.class_obj = self._make_class(self.teacher_profile, code='BSIT-1A', subject=subject)
        Enrollment.objects.create(
            student_profile=self.student, class_obj=self.class_obj,
            status=Enrollment.STATUS_ACTIVE,
        )

        today = timezone.localdate()
        slot = self._make_slot(self.class_obj, weekday=0)
        statuses = [
            (ExcuseSlip.STATUS_PENDING, 'Fever'),
            (ExcuseSlip.STATUS_APPROVED, 'Family emergency'),
            (ExcuseSlip.STATUS_REJECTED, 'No documentation'),
        ]
        for i, (status, reason) in enumerate(statuses):
            session = self._make_session(slot, today - timedelta(days=10 + i))
            self._make_record(session, self.student, AttendanceRecord.STATUS_ABSENT)
            ExcuseSlip.objects.create(
                student_profile=self.student, session=session,
                reason=reason, status=status,
            )

        self.client.force_login(self.student_user)

    def test_dashboard_renders_excuse_slip_tracker(self):
        res = self.client.get(reverse('student_dashboard'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'My Excuse Slips')
        self.assertContains(res, 'Pending Review')
        self.assertContains(res, 'Fever')

    def test_dashboard_stats_expose_slip_counts(self):
        stats = _student_dashboard_stats(self.student)
        self.assertEqual(stats['excuse_slip_counts']['pending'], 1)
        self.assertEqual(stats['excuse_slip_counts']['approved'], 1)
        self.assertEqual(stats['excuse_slip_counts']['rejected'], 1)
        self.assertEqual(len(stats['excuse_slips']), 3)


# // ==============================================================================
# // S-06: FILTERABLE MEETING-BY-MEETING ATTENDANCE LOG
# // ==============================================================================
class AttendanceLogFilterTests(TestCase, BaseWeek5FixtureMixin):
    def setUp(self):
        self.department = self._make_department()
        self.program = self._make_program()
        self.academic_year = self._make_academic_year()
        self.teacher_user, self.teacher_profile = self._make_teacher('w5_s6_teacher')
        self.student_user, self.student = self._make_student('w5_s6', '2026-9301')

        subject_a = Subject.objects.create(code='IT101', name='Intro to Computing')
        subject_b = Subject.objects.create(code='IT102', name='Programming 1')
        self.class_a = self._make_class(self.teacher_profile, code='BSIT-1A', subject=subject_a)
        self.class_b = self._make_class(self.teacher_profile, code='BSIT-1B', subject=subject_b)
        for class_obj in (self.class_a, self.class_b):
            Enrollment.objects.create(
                student_profile=self.student, class_obj=class_obj,
                status=Enrollment.STATUS_ACTIVE,
            )

        today = timezone.localdate()
        slot_a = self._make_slot(self.class_a, weekday=0)
        slot_b = self._make_slot(self.class_b, weekday=1)
        session_a = self._make_session(slot_a, today - timedelta(days=7))
        session_b = self._make_session(slot_b, today - timedelta(days=6))
        self._make_record(session_a, self.student, AttendanceRecord.STATUS_PRESENT)
        self._make_record(session_b, self.student, AttendanceRecord.STATUS_ABSENT)

        self.client.force_login(self.student_user)

    def test_log_filters_by_status(self):
        res = self.client.get(reverse('student_attendance_log'), {'status': 'ABSENT'})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.context['selected_status'], 'ABSENT')
        records = list(res.context['records'])
        self.assertTrue(records)
        self.assertTrue(all(r.status == AttendanceRecord.STATUS_ABSENT for r in records))
        self.assertTrue(all(
            r.session.slot.class_obj.class_code == 'BSIT-1B' for r in records))

    def test_log_filters_by_class(self):
        res = self.client.get(
            reverse('student_attendance_log'), {'class_id': self.class_a.pk},
        )
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.context['selected_class_id'], self.class_a.pk)
        records = list(res.context['records'])
        self.assertTrue(records)
        self.assertTrue(all(
            r.session.slot.class_obj.class_code == 'BSIT-1A' for r in records))

    def test_log_renders_filter_controls_and_slot_time(self):
        res = self.client.get(reverse('student_attendance_log'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'name="status"')
        self.assertContains(res, 'name="class_id"')
        self.assertContains(res, '08:00 AM')  # slot window column


# // ==============================================================================
# // A-10: CSV IMPORT 3-STEP PREVIEW
# // ==============================================================================
CSV_FIXTURE = (
    "student_number,first_name,last_name,email,program,year_level,section\n"
    "2027-9001,Ana,Cruz,ana9001@kcc.edu.ph,BSIT,1,BSIT-1A\n"
    "2027-9002,Ben,Reyes,ben9002@kcc.edu.ph,BSIT,2,\n"
    "2027-9003,Cat,Santos,cat9003@kcc.edu.ph,NOPE,1,\n"
)


class CsvImportPreviewTests(TestCase, BaseWeek5FixtureMixin):
    def setUp(self):
        self.department = self._make_department()
        self.program = self._make_program()
        self.academic_year = self._make_academic_year()
        self.admin = self._make_admin()
        self.client.force_login(self.admin)

    def _upload(self):
        csv_file = SimpleUploadedFile(
            'roster.csv', CSV_FIXTURE.encode('utf-8'), content_type='text/csv',
        )
        return self.client.post(reverse('student_directory_import_csv'), {
            'csv_file': csv_file,
            'auto_approve': 'on',
            'preview': '1',
        })

    def test_preview_validates_without_creating_anything(self):
        res = self._upload()
        self.assertEqual(res.status_code, 200)
        preview = res.context['preview']
        self.assertIsNotNone(preview)
        self.assertEqual(preview['valid_count'], 2)
        self.assertEqual(preview['invalid_count'], 1)
        self.assertEqual(len(preview['rows']), 3)
        bad = [r for r in preview['rows'] if not r['ok']]
        self.assertEqual(len(bad), 1)
        self.assertIn('NOPE', bad[0]['error'])
        self.assertEqual(
            Profile.objects.filter(student_number__startswith='2027-9').count(), 0,
        )

    def test_confirm_imports_previewed_rows(self):
        res = self.client.post(reverse('student_directory_import_csv'), {
            'confirm': '1',
            'csv_text': CSV_FIXTURE,
            'auto_approve': 'on',
            'default_password': 'KccStudent2026!',
        })
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            Profile.objects.filter(student_number__startswith='2027-9').count(), 2,
        )
        created = Profile.objects.get(student_number='2027-9001')
        self.assertEqual(created.approval_status, Profile.APPROVAL_APPROVED)
        self.assertEqual(len(created.qr_secret), 64)
        self.assertFalse(Profile.objects.filter(student_number='2027-9003').exists())

    def test_direct_upload_without_preview_still_imports(self):
        csv_file = SimpleUploadedFile(
            'roster.csv', CSV_FIXTURE.encode('utf-8'), content_type='text/csv',
        )
        res = self.client.post(reverse('student_directory_import_csv'), {
            'csv_file': csv_file,
            'auto_approve': 'on',
        })
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            Profile.objects.filter(student_number__startswith='2027-9').count(), 2,
        )

    def test_confirm_without_csv_text_redirects(self):
        res = self.client.post(reverse('student_directory_import_csv'), {
            'confirm': '1',
            'csv_text': '',
        })
        self.assertEqual(res.status_code, 302)
        self.assertEqual(res.url, reverse('student_directory_import_csv'))


# // ==============================================================================
# // T-06: MIDTERM / FINALS SPLIT EXPORT
# // ==============================================================================
class SplitExportTests(TestCase, BaseWeek5FixtureMixin):
    def setUp(self):
        self.department = self._make_department()
        self.program = self._make_program()
        self.academic_year = self._make_academic_year()
        self.teacher_user, self.teacher_profile = self._make_teacher('w5_t6_teacher')
        subject = Subject.objects.create(code='IT101', name='Introduction to Computing')
        self.class_obj = self._make_class(self.teacher_profile, code='BSIT-1A', subject=subject)

        today = timezone.localdate()
        slot = self._make_slot(self.class_obj, weekday=0)
        for i in range(20):
            self._make_session(slot, today - timedelta(days=60 - i))

        self.client.force_login(self.teacher_user)

    def _week_columns(self, res):
        wb = load_workbook(BytesIO(res.content))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        return ws.title, [h for h in headers if h and str(h).startswith('Week ')]

    def test_midterm_export_caps_to_first_9_meetings(self):
        res = self.client.get(
            reverse('teacher_class_term_export', args=[self.class_obj.pk]),
            {'range': 'midterm'},
        )
        self.assertEqual(res.status_code, 200)
        title, weeks = self._week_columns(res)
        self.assertEqual(title, 'Midterm Attendance')
        self.assertEqual(len(weeks), 9)
        self.assertIn('midterm_attendance_BSIT-1A.xlsx', res['Content-Disposition'])

    def test_finals_export_caps_to_first_18_meetings(self):
        res = self.client.get(
            reverse('teacher_class_term_export', args=[self.class_obj.pk]),
            {'range': 'finals'},
        )
        self.assertEqual(res.status_code, 200)
        title, weeks = self._week_columns(res)
        self.assertEqual(title, 'Finals Attendance')
        self.assertEqual(len(weeks), 18)
        self.assertIn('finals_attendance_BSIT-1A.xlsx', res['Content-Disposition'])

    def test_term_export_includes_all_meetings(self):
        res = self.client.get(
            reverse('teacher_class_term_export', args=[self.class_obj.pk]),
        )
        self.assertEqual(res.status_code, 200)
        title, weeks = self._week_columns(res)
        self.assertEqual(title, 'Term Attendance')
        self.assertEqual(len(weeks), 20)
