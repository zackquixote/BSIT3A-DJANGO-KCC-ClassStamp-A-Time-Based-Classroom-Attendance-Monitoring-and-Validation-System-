"""
# // ==============================================================================
# // KCC CLASSSTAMP - WEEK 3 REPORTS & ANALYTICS TESTS
# // ==============================================================================
# // Covers:
# //   - A-09: Admin audit log actor filter (page + CSV export)
# //   - S-04: CHED 20% 4-stage early-warning sentinel (student dashboard)
# //   - T-04: Registrar term-wide Excel (.xlsx) export per class section
# //   - S-05: Visual weekly timetable grid (Mon-Sat, active slots only)
# //   - T-05: Consecutive 3-absence sentinel on teacher dashboard
# // ==============================================================================
"""

from datetime import date, time, timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.models import (
    AcademicYear,
    AttendanceRecord,
    AuditEvent,
    Class,
    ClassSession,
    ClassSlot,
    Department,
    Enrollment,
    Profile,
    Program,
    Subject,
)
from apps.core.views.helpers.dashboard_helpers import (
    _student_dashboard_stats,
    _teacher_dashboard_stats,
)

User = get_user_model()


class BaseDashboardFixtureMixin:
    def _make_department(self):
        return Department.objects.create(code='CCS', name='College of Computer Studies')

    def _make_program(self):
        return Program.objects.create(
            department=self.department, code='BSIT',
            name='BS Information Technology', years_total=4,
        )

    def _make_academic_year(self):
        return AcademicYear.objects.create(
            name='2025-2026', starts_on=date(2025, 6, 1), ends_on=date(2026, 3, 31), is_active=True,
        )

    def _make_teacher(self, username):
        user = User.objects.create_user(username=username, password='Teacher!Pass123')
        profile = Profile.objects.create(
            user=user, role=Profile.ROLE_TEACHER, department=self.department,
        )
        return user, profile

    def _make_student(self, username, student_number, first_name='Ana', last_name='Cruz'):
        user = User.objects.create_user(
            username=username, password='Student!Pass123',
            first_name=first_name, last_name=last_name,
        )
        profile = Profile.objects.create(
            user=user, role=Profile.ROLE_STUDENT, student_number=student_number,
            program=self.program, year_level=1,
            approval_status=Profile.APPROVAL_APPROVED, qr_secret='a' * 64,
        )
        return user, profile

    def _make_class(self, teacher_profile, subject, code='BSIT-1A'):
        class_obj = Class.objects.create(
            academic_year=self.academic_year, class_code=code,
            subject=subject, teacher_profile=teacher_profile, grace_minutes=10,
        )
        return class_obj

    def _make_slot(self, class_obj, weekday, starts_at, ends_at, room='Room 1', is_active=True):
        return ClassSlot.objects.create(
            class_obj=class_obj, weekday=weekday,
            starts_at=starts_at, ends_at=ends_at, room=room, is_active=is_active,
        )

    def _make_session(self, slot, session_date, status=ClassSession.STATUS_CLOSED):
        return ClassSession.objects.create(
            slot=slot, session_date=session_date, status=status,
        )

    def _make_record(self, session, student_profile, status):
        return AttendanceRecord.objects.create(
            session=session, student_profile=student_profile,
            status=status, method=AttendanceRecord.METHOD_SYSTEM,
            stamped_at=timezone.now(),
        )


# // ==============================================================================
# // A-09: ADMIN AUDIT LOG ACTOR FILTER
# // ==============================================================================
class AuditActorFilterTests(TestCase):
    def setUp(self):
        self.admin_user = User.objects.create_user(
            username='audit_admin', password='Admin!Pass123', first_name='Sys', last_name='Admin',
        )
        Profile.objects.create(user=self.admin_user, role=Profile.ROLE_ADMIN)

        self.actor_one = User.objects.create_user(
            username='admin_one', password='Admin!Pass123', first_name='Ada', last_name='One',
        )
        self.actor_two = User.objects.create_user(
            username='admin_two', password='Admin!Pass123', first_name='Bob', last_name='Two',
        )

        AuditEvent.objects.create(
            actor=self.actor_one, action='APPROVE_STUDENT',
            object_type='Profile', object_id='OBJ-ONE-123',
        )
        AuditEvent.objects.create(
            actor=self.actor_one, action='ROTATE_QR',
            object_type='Profile', object_id='OBJ-ONE-124',
        )
        AuditEvent.objects.create(
            actor=self.actor_two, action='REJECT_STUDENT',
            object_type='Profile', object_id='OBJ-TWO-456',
        )

        self.client.force_login(self.admin_user)

    def test_audit_log_renders_actor_dropdown(self):
        res = self.client.get(reverse('admin_audit_log'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'name="actor"')
        self.assertContains(res, 'admin_one')
        self.assertContains(res, 'admin_two')
        # 3 fixtures + 1 USER_LOGIN event emitted by force_login()
        self.assertEqual(res.context['total_count'], 4)

    def test_audit_log_filters_by_actor(self):
        res = self.client.get(reverse('admin_audit_log'), {'actor': 'admin_one'})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.context['total_count'], 2)
        self.assertEqual(res.context['actor_filter'], 'admin_one')
        self.assertContains(res, 'OBJ-ONE-123')
        self.assertContains(res, 'OBJ-ONE-124')
        self.assertNotContains(res, 'OBJ-TWO-456')

    def test_csv_export_respects_actor_filter(self):
        res = self.client.get(
            reverse('admin_audit_log_export_csv'), {'actor': 'admin_one'},
        )
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'admin_one')
        self.assertContains(res, 'Ada One')
        self.assertNotContains(res, 'admin_two')
        self.assertNotContains(res, 'OBJ-TWO-456')

    def test_audit_log_requires_admin_role(self):
        self.client.logout()
        intruder = User.objects.create_user(username='intruder', password='Admin!Pass123')
        Profile.objects.create(user=intruder, role=Profile.ROLE_TEACHER, department=Department.objects.create(code='X', name='X'))
        self.client.force_login(intruder)
        res = self.client.get(reverse('admin_audit_log'))
        self.assertEqual(res.status_code, 302)


# // ==============================================================================
# // S-04: CHED 20% EARLY-WARNING SENTINEL (4 STAGES)
# // ==============================================================================
class StudentSentinelStageTests(TestCase, BaseDashboardFixtureMixin):
    def setUp(self):
        self.department = self._make_department()
        self.program = self._make_program()
        self.academic_year = self._make_academic_year()
        self.teacher_user, self.teacher_profile = self._make_teacher('sentinel_teacher')
        self.student_user, self.student_profile = self._make_student(
            'sentinel_student', '2026-6101',
        )

        today = timezone.localdate()
        # (code, subject, specs) -> specs are per-session statuses
        self.classes = {}
        scenarios = [
            ('GREEN', 'IT101', ['PRESENT', 'PRESENT', 'PRESENT', 'PRESENT']),
            ('AMBER', 'IT102', ['ABSENT'] + ['PRESENT'] * 9),
            ('ORANGE', 'IT103', ['ABSENT', 'EXCUSED'] + ['PRESENT'] * 4),
            ('RED', 'IT104', ['ABSENT', 'ABSENT', 'PRESENT', 'PRESENT', 'PRESENT']),
        ]
        for idx, (tag, subject_code, specs) in enumerate(scenarios):
            subject = Subject.objects.create(code=subject_code, name=f'Subject {tag}')
            class_obj = self._make_class(self.teacher_profile, subject, code=f'BSIT-{tag}')
            Enrollment.objects.create(
                student_profile=self.student_profile, class_obj=class_obj,
                status=Enrollment.STATUS_ACTIVE,
            )
            slot = self._make_slot(class_obj, weekday=0, starts_at=time(8, 0), ends_at=time(9, 0))
            for i, status in enumerate(specs):
                session = self._make_session(slot, today - timedelta(days=40 - i * 2))
                self._make_record(session, self.student_profile, status)
            self.classes[tag] = class_obj

        self.client.force_login(self.student_user)

    def _breakdown_by_tag(self):
        stats = _student_dashboard_stats(self.student_profile)
        by_class = {b['class_obj'].id: b for b in stats['class_breakdowns']}
        return {tag: by_class[self.classes[tag].id] for tag in self.classes}

    def test_sentinel_stages_are_classified_correctly(self):
        breakdowns = self._breakdown_by_tag()
        self.assertEqual(breakdowns['GREEN']['sentinel']['stage'], 'green')
        self.assertEqual(breakdowns['GREEN']['sentinel']['label'], 'Excellent Attendance')
        self.assertEqual(breakdowns['GREEN']['absence_rate'], 0.0)

        self.assertEqual(breakdowns['AMBER']['sentinel']['stage'], 'amber')
        self.assertEqual(breakdowns['AMBER']['sentinel']['label'], 'Caution')
        self.assertEqual(breakdowns['AMBER']['absence_rate'], 10.0)

        self.assertEqual(breakdowns['ORANGE']['sentinel']['stage'], 'orange')
        self.assertEqual(breakdowns['ORANGE']['sentinel']['label'], 'Warning')
        self.assertEqual(breakdowns['ORANGE']['absence_rate'], 16.7)

        self.assertEqual(breakdowns['RED']['sentinel']['stage'], 'red')
        self.assertEqual(breakdowns['RED']['sentinel']['label'], 'Exceeded threshold — Contact Instructor/Dean')
        self.assertEqual(breakdowns['RED']['absence_rate'], 40.0)

    def test_dashboard_renders_sentinel_banners(self):
        res = self.client.get(reverse('student_dashboard'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'Excellent Attendance')
        self.assertContains(res, 'Caution')
        self.assertContains(res, 'Warning')
        self.assertContains(res, 'Exceeded threshold — Contact Instructor/Dean')

    def test_no_meetings_class_has_no_sentinel_warning(self):
        subject = Subject.objects.create(code='IT105', name='Empty Class')
        class_obj = self._make_class(self.teacher_profile, subject, code='BSIT-EMPTY')
        Enrollment.objects.create(
            student_profile=self.student_profile, class_obj=class_obj,
            status=Enrollment.STATUS_ACTIVE,
        )
        breakdowns = _student_dashboard_stats(self.student_profile)['class_breakdowns']
        empty = next(b for b in breakdowns if b['class_obj'].id == class_obj.id)
        self.assertEqual(empty['sentinel']['stage'], 'none')
        self.assertEqual(empty['sentinel']['label'], 'No meetings yet')
        self.assertIsNone(empty['absence_rate'])


# // ==============================================================================
# // T-04: REGISTRAR TERM-WIDE EXCEL EXPORT
# // ==============================================================================
class TermExcelExportTests(TestCase, BaseDashboardFixtureMixin):
    def setUp(self):
        self.department = self._make_department()
        self.program = self._make_program()
        self.academic_year = self._make_academic_year()
        self.subject = Subject.objects.create(code='IT311', name='Capstone 1')

        self.teacher_user, self.teacher_profile = self._make_teacher('export_teacher')
        self.other_teacher_user, self.other_teacher_profile = self._make_teacher('export_other')
        self.class_obj = self._make_class(self.teacher_profile, self.subject, code='BSIT-3A')

        self.student_a_user, self.student_a = self._make_student(
            'export_a', '2026-9001', first_name='Alisa', last_name='Alpha',
        )
        self.student_b_user, self.student_b = self._make_student(
            'export_b', '2026-9002', first_name='Ben', last_name='Bravo',
        )
        for student in (self.student_a, self.student_b):
            Enrollment.objects.create(
                student_profile=student, class_obj=self.class_obj,
                status=Enrollment.STATUS_ACTIVE,
            )

        today = timezone.localdate()
        slot = self._make_slot(self.class_obj, weekday=0, starts_at=time(8, 0), ends_at=time(9, 0))
        self.sessions = [
            self._make_session(slot, today - timedelta(days=20)),
            self._make_session(slot, today - timedelta(days=13)),
        ]
        for session in self.sessions:
            self._make_record(session, self.student_a, AttendanceRecord.STATUS_PRESENT)
            self._make_record(session, self.student_b, AttendanceRecord.STATUS_ABSENT)

    def test_export_returns_valid_xlsx_with_totals_and_red_flag(self):
        self.client.force_login(self.teacher_user)
        res = self.client.get(reverse('teacher_class_term_export', args=[self.class_obj.pk]))
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res['Content-Type'].startswith(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))
        self.assertIn('BSIT-3A', res['Content-Disposition'])

        wb = load_workbook(BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws.title, 'Term Attendance')
        self.assertEqual(ws.freeze_panes, 'F2')

        headers = [c.value for c in ws[1]]
        self.assertEqual(headers[:5], ['Student ID', 'Last Name', 'First Name', 'Program', 'Year'])
        self.assertEqual(headers[5], f"Week 1 ({self.sessions[0].session_date.strftime('%m/%d')})")
        self.assertEqual(headers[6], f"Week 2 ({self.sessions[1].session_date.strftime('%m/%d')})")
        self.assertEqual(headers[7:], ['Total Present', 'Total Late', 'Total Absent', 'Attendance %'])

        # Row 2: Alisa Alpha — all present
        self.assertEqual(ws.cell(row=2, column=1).value, '2026-9001')
        self.assertEqual(ws.cell(row=2, column=6).value, 'P')
        self.assertEqual(ws.cell(row=2, column=8).value, 2)
        self.assertEqual(ws.cell(row=2, column=11).value, 100.0)
        self.assertFalse(str(ws.cell(row=2, column=1).fill.fgColor.rgb).endswith('FECACA'))

        # Row 3: Ben Bravo — all absent -> 0% -> red CHED flag
        self.assertEqual(ws.cell(row=3, column=1).value, '2026-9002')
        self.assertEqual(ws.cell(row=3, column=6).value, 'A')
        self.assertEqual(ws.cell(row=3, column=10).value, 2)
        self.assertEqual(ws.cell(row=3, column=11).value, 0.0)
        self.assertTrue(str(ws.cell(row=3, column=1).fill.fgColor.rgb).endswith('FECACA'))
        self.assertTrue(ws.cell(row=3, column=11).font.bold)

    def test_export_forbidden_for_unassigned_teacher(self):
        self.client.force_login(self.other_teacher_user)
        res = self.client.get(reverse('teacher_class_term_export', args=[self.class_obj.pk]))
        self.assertEqual(res.status_code, 302)
        self.assertEqual(res.url, reverse('teacher_dashboard'))

    def test_export_requires_login(self):
        res = self.client.get(reverse('teacher_class_term_export', args=[self.class_obj.pk]))
        self.assertEqual(res.status_code, 302)


# // ==============================================================================
# // S-05: VISUAL WEEKLY TIMETABLE GRID
# // ==============================================================================
class StudentTimetableGridTests(TestCase, BaseDashboardFixtureMixin):
    def setUp(self):
        self.department = self._make_department()
        self.program = self._make_program()
        self.academic_year = self._make_academic_year()
        self.teacher_user, self.teacher_profile = self._make_teacher('grid_teacher')
        self.student_user, self.student_profile = self._make_student('grid_student', '2026-6201')

        self.subject = Subject.objects.create(code='IT121', name='Web Systems')
        self.class_obj = self._make_class(self.teacher_profile, self.subject, code='BSIT-2A')
        Enrollment.objects.create(
            student_profile=self.student_profile, class_obj=self.class_obj,
            status=Enrollment.STATUS_ACTIVE,
        )
        self._make_slot(self.class_obj, weekday=0, starts_at=time(8, 0), ends_at=time(9, 0), room='Room 101')
        self._make_slot(self.class_obj, weekday=2, starts_at=time(13, 0), ends_at=time(14, 0), room='Lab 3')
        self._make_slot(self.class_obj, weekday=6, starts_at=time(7, 0), ends_at=time(8, 0), room='SUNDAY-ROOM')
        self._make_slot(
            self.class_obj, weekday=1, starts_at=time(10, 0), ends_at=time(11, 0),
            room='INACTIVE-ROOM', is_active=False,
        )

    def test_helper_builds_filtered_sorted_timetable(self):
        stats = _student_dashboard_stats(self.student_profile)
        slots = stats['timetable_slots']
        self.assertEqual(len(slots), 2)

        mon, wed = slots
        self.assertEqual(mon['weekday'], 0)
        self.assertEqual(mon['weekday_label'], 'Monday')
        self.assertEqual(mon['starts_at'], time(8, 0))
        self.assertEqual(mon['room'], 'Room 101')
        self.assertEqual(mon['class_code'], 'BSIT-2A')
        self.assertEqual(mon['subject_code'], 'IT121')
        self.assertEqual(mon['teacher'], 'grid_teacher')
        self.assertTrue(mon['color'].startswith('#'))
        self.assertEqual(wed['weekday'], 2)
        self.assertEqual(wed['room'], 'Lab 3')

        weekdays = {s['weekday'] for s in slots}
        self.assertNotIn(6, weekdays)
        rooms = {s['room'] for s in slots}
        self.assertNotIn('SUNDAY-ROOM', rooms)
        self.assertNotIn('INACTIVE-ROOM', rooms)
        self.assertEqual(stats['today_weekday'], timezone.localdate().weekday())

    def test_dashboard_renders_grid_section(self):
        self.client.force_login(self.student_user)
        res = self.client.get(reverse('student_dashboard'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'Weekly Class Grid')


# // ==============================================================================
# // T-05: CONSECUTIVE 3-ABSENCE SENTINEL
# // ==============================================================================
class ConsecutiveAbsenceSentinelTests(TestCase, BaseDashboardFixtureMixin):
    def setUp(self):
        self.department = self._make_department()
        self.program = self._make_program()
        self.academic_year = self._make_academic_year()
        self.subject = Subject.objects.create(code='IT122', name='Data Analytics')

        self.teacher_user, self.teacher_profile = self._make_teacher('streak_teacher')
        self.class_obj = self._make_class(self.teacher_profile, self.subject, code='BSIT-2B')

        self.streak_user, self.streak_student = self._make_student(
            'streak_a', '2026-6301', first_name='Ana', last_name='Cruz',
        )
        self.clean_user, self.clean_student = self._make_student(
            'streak_b', '2026-6302', first_name='Ben', last_name='Doe',
        )
        for student in (self.streak_student, self.clean_student):
            Enrollment.objects.create(
                student_profile=student, class_obj=self.class_obj,
                status=Enrollment.STATUS_ACTIVE,
            )

        today = timezone.localdate()
        slot = self._make_slot(self.class_obj, weekday=0, starts_at=time(8, 0), ends_at=time(9, 0))
        # Ana: 3 newest meetings ABSENT (streak of 3), older meeting PRESENT
        # Ben: 2 ABSENT then PRESENT -> no flag
        specs = {
            self.streak_student: ['ABSENT', 'ABSENT', 'ABSENT', 'PRESENT'],
            self.clean_student: ['ABSENT', 'ABSENT', 'PRESENT'],
        }
        for i in range(4):
            session = self._make_session(slot, today - timedelta(days=i + 1))
            for student, statuses in specs.items():
                if i < len(statuses):
                    self._make_record(session, student, statuses[i])

    def test_helper_flags_three_absent_streak(self):
        stats = _teacher_dashboard_stats(self.teacher_profile)
        classes = list(stats['classes'])
        self.assertEqual(len(classes), 1)
        flagged = classes[0].absent_streaks
        self.assertEqual(len(flagged), 1)
        self.assertEqual(flagged[0]['first_name'], 'Ana')
        self.assertEqual(flagged[0]['last_name'], 'Cruz')
        self.assertEqual(flagged[0]['student_number'], '2026-6301')
        self.assertEqual(flagged[0]['streak'], 3)
        self.assertEqual(stats['absent_streak_students'], 1)

    def test_two_absent_streak_not_flagged(self):
        stats = _teacher_dashboard_stats(self.teacher_profile)
        flagged = list(stats['classes'])[0].absent_streaks
        names = {(f['first_name'], f['last_name']) for f in flagged}
        self.assertNotIn(('Ben', 'Doe'), names)

    def test_dashboard_renders_streak_chips(self):
        self.client.force_login(self.teacher_user)
        res = self.client.get(reverse('teacher_dashboard'))
        self.assertEqual(res.status_code, 200)
        self.assertContains(res, 'Ana Cruz')
        self.assertContains(res, '2026-6301')
        self.assertContains(res, '3 straight absences')
        self.assertNotContains(res, '2 straight absences')
